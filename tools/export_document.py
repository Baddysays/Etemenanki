#!/usr/bin/env python3
"""Export translated documents for Etemenanki hub workflows."""
from __future__ import annotations

import json
import re
import sys
from copy import deepcopy
from pathlib import Path


def emit(payload: dict) -> None:
    sys.stdout.buffer.write(json.dumps(payload, ensure_ascii=True).encode("utf-8"))
    sys.stdout.buffer.write(b"\n")
    sys.stdout.buffer.flush()


def _segments(meta: dict) -> list[dict]:
    return list(meta.get("segments") or [])


def export_subtitle(meta: dict, output_path: Path) -> None:
    fmt = str(meta.get("format") or "srt").lower()
    segments = _segments(meta)
    lines: list[str] = []
    if fmt == "ass":
        lines.append("[Script Info]")
        lines.append("Title: Translated")
        lines.append("")
        lines.append("[Events]")
        lines.append("Format: Layer, Start, End, Style, Name, MarginL, MarginR, MarginV, Effect, Text")
        for item in segments:
            text = item.get("text_translated") or item.get("text") or ""
            if not text:
                continue
            start = item.get("start", "0:00:00.00")
            end = item.get("end", "0:00:01.00")
            lines.append(f"Dialogue: 0,{start},{end},Default,,0,0,0,,{text}")
        output_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
        return

    if fmt == "vtt":
        lines.append("WEBVTT")
        lines.append("")
        for item in segments:
            text = item.get("text_translated") or item.get("text") or ""
            if not text:
                continue
            start = str(item.get("start", "00:00:00.000")).replace(",", ".")
            end = str(item.get("end", "00:00:01.000")).replace(",", ".")
            lines.extend(["", f"{start} --> {end}", text])
        output_path.write_text("\n".join(lines).strip() + "\n", encoding="utf-8")
        return

    for idx, item in enumerate(segments, start=1):
        text = item.get("text_translated") or item.get("text") or ""
        if not text:
            continue
        start = item.get("start", "00:00:00,000")
        end = item.get("end", "00:00:01,000")
        lines.extend([str(idx), f"{start} --> {end}", text, ""])
    output_path.write_text("\n".join(lines).strip() + "\n", encoding="utf-8")


def export_docx(source_path: Path, meta: dict, output_path: Path) -> None:
    from docx import Document

    doc = Document(str(source_path))
    mapping = {
        str(item.get("id")): item.get("text_translated") or item.get("text") or ""
        for item in _segments(meta)
        if item.get("id")
    }
    for idx, paragraph in enumerate(doc.paragraphs):
        key = f"p{idx}"
        if key in mapping and mapping[key].strip():
            paragraph.text = mapping[key]
    doc.save(str(output_path))


def export_spreadsheet(source_path: Path, meta: dict, output_path: Path) -> None:
    fmt = str(meta.get("format") or source_path.suffix.lower().lstrip(".")).lower()
    mapping = {
        str(item.get("id")): item.get("text_translated") or item.get("text") or ""
        for item in _segments(meta)
        if item.get("id")
    }

    if fmt == "csv":
        import csv

        rows: list[list[str]] = []
        with source_path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.reader(handle)
            for r_idx, row in enumerate(reader):
                updated = []
                for c_idx, value in enumerate(row):
                    key = f"r{r_idx}c{c_idx}"
                    updated.append(mapping.get(key, value))
                rows.append(updated)
        with output_path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerows(rows)
        return

    from openpyxl import load_workbook

    wb = load_workbook(str(source_path))
    for item in _segments(meta):
        cell_ref = item.get("cell")
        sheet_name = item.get("sheet")
        translated = item.get("text_translated") or item.get("text")
        if not cell_ref or not sheet_name or not translated:
            continue
        wb[sheet_name][cell_ref].value = translated
    wb.save(str(output_path))


def export_json(meta: dict, output_path: Path) -> None:
    root = deepcopy(meta.get("root") or {})
    mapping = {
        str(item.get("id")): item.get("text_translated") or item.get("text") or ""
        for item in _segments(meta)
        if item.get("id")
    }

    def set_path(obj, path_key: str, value: str) -> None:
        if not path_key:
            return
        if path_key.startswith("[") and path_key.endswith("]"):
            return
        parts = re.split(r"\.(?![^\[]*\])", path_key)
        current = obj
        for part in parts[:-1]:
            if "[" in part:
                name, idx = part.split("[", 1)
                idx = int(idx.rstrip("]"))
                current = current[name][idx]
            else:
                current = current[part]
        last = parts[-1]
        if "[" in last:
            name, idx = last.split("[", 1)
            idx = int(idx.rstrip("]"))
            current[name][idx] = value
        else:
            current[last] = value

    for key, value in mapping.items():
        set_path(root, key, value)

    output_path.write_text(json.dumps(root, ensure_ascii=False, indent=2), encoding="utf-8")


def export_html(meta: dict, output_path: Path) -> None:
    raw = str(meta.get("raw_html") or "")
    translated = ""
    for item in _segments(meta):
        translated = item.get("text_translated") or item.get("text") or translated
    if raw and translated:
        import html as _html
        safe = _html.escape(translated)
        body = re.sub(r"(?is)<body[^>]*>.*?</body>", f"<body><pre>{safe}</pre></body>", raw, count=1)
        if body != raw:
            output_path.write_text(body, encoding="utf-8")
            return
    output_path.write_text(translated or raw, encoding="utf-8")


def export_epub(source_path: Path, meta: dict, output_path: Path) -> None:
    try:
        from ebooklib import epub
    except ImportError as exc:
        raise RuntimeError("Install ebooklib: pip install ebooklib") from exc

    book = epub.read_epub(str(source_path))
    mapping = {
        str(item.get("id")): item.get("text_translated") or item.get("text") or ""
        for item in _segments(meta)
        if item.get("id")
    }
    idx = 0
    for item in book.get_items():
        if item.get_type() != 9:
            continue
        key = f"ch{idx}"
        if key in mapping:
            import html as _html
            text = _html.escape(mapping[key])
            html_body = f"<?xml version='1.0' encoding='utf-8'?><html><body><p>{text}</p></body></html>"
            item.set_content(html_body.encode("utf-8"))
        idx += 1
    epub.write_epub(str(output_path), book)


def main() -> int:
    if len(sys.argv) < 2:
        emit({"ok": False, "error": "usage: export_document.py <meta.json>"})
        return 2

    meta_path = Path(sys.argv[1])
    if not meta_path.is_file():
        emit({"ok": False, "error": f"meta file not found: {meta_path}"})
        return 2

    envelope = json.loads(meta_path.read_text(encoding="utf-8"))
    workflow_id = str(envelope.get("workflow_id") or "")
    meta = dict(envelope.get("workflow_meta") or {})
    source_path = Path(str(envelope.get("source_path") or ""))
    output_path = Path(str(envelope.get("output_path") or ""))
    output_path.parent.mkdir(parents=True, exist_ok=True)

    try:
        if workflow_id == "subtitle":
            export_subtitle(meta, output_path)
        elif workflow_id == "docx":
            export_docx(source_path, meta, output_path)
        elif workflow_id == "spreadsheet":
            export_spreadsheet(source_path, meta, output_path)
        elif workflow_id == "json":
            export_json(meta, output_path)
        elif workflow_id == "html":
            export_html(meta, output_path)
        elif workflow_id == "epub":
            export_epub(source_path, meta, output_path)
        else:
            text = "\n\n".join(
                item.get("text_translated") or item.get("text") or "" for item in _segments(meta)
            ).strip()
            if not text:
                emit({"ok": False, "error": f"unsupported export workflow: {workflow_id}"})
                return 2
            output_path.write_text(text, encoding="utf-8")

        emit({"ok": True, "output_path": str(output_path)})
        return 0
    except Exception as exc:  # noqa: BLE001
        emit({"ok": False, "error": str(exc) or repr(exc)})
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
